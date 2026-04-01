import pandas as pd
from openpyxl import load_workbook

def run(input_path, output_path):
    df = pd.read_excel(
        input_path,
        header=1,
        dtype={
            3: str,
            7: str,
            37: str
        }
    )

    cols = df.columns.tolist()
    辅助单号 = next(c for c in cols if "辅助单号" in str(c))
    组合产品 = next(c for c in cols if "组合产品" in str(c))
    子订单本地 = next(c for c in cols if "子订单本地" in str(c))

    df_final = df.groupby(辅助单号, as_index=False).agg(
        **{
            组合产品: (组合产品, lambda x: '是' if '是' in x.values else '否'),
            子订单本地: (子订单本地, 'sum'),
            **{c: (c, 'first') for c in cols if c not in [辅助单号, 组合产品, 子订单本地]}
        }
    )
    df_final = df_final[cols]

    wb = load_workbook(input_path)
    ws = wb.active
    ws.delete_rows(3, ws.max_row)

    for row_idx, row in df_final.iterrows():
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx + 3, column=col_idx, value=val)

    ws.column_dimensions['D'].number_format = '@'
    ws.column_dimensions['H'].number_format = '@'
    ws.column_dimensions['AL'].number_format = '@'
    ws.column_dimensions['AR'].number_format = 'yyyy-mm-dd hh:mm:ss'
    ws.column_dimensions['AS'].number_format = 'yyyy-mm-dd hh:mm:ss'

    for col in ['D', 'H', 'AL', 'AR', 'AS']:
        ws.column_dimensions[col].width = 20

    wb.save(output_path)